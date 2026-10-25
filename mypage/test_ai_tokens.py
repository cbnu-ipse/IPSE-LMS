import io
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings

from core.models import AIUsage

from . import ai
from .models import GeneratedQuestion, PersonalDocument, ProcessingStatus
from .views import _generate_question_bg

User = get_user_model()


class ExtractionTestCase(TestCase):
    def test_clean_markdown_removes_noise(self):
        text = "강의 머리글\n본문 1\n\n\n\n3\n강의 머리글\n본문 2\n강의 머리글\n![그림](a.png)\n|  a  |  b  |\n|---|---|\n- 12 -"
        cleaned = ai.clean_markdown(text)
        self.assertNotIn("강의 머리글", cleaned)
        self.assertNotIn("![", cleaned)
        self.assertNotIn("\n\n\n", cleaned)
        self.assertIn("본문 1", cleaned)
        self.assertIn("|a|b|", cleaned)

    def test_docx_and_pptx_become_markdown(self):
        import docx
        from pptx import Presentation

        d = docx.Document()
        d.add_heading("정규화", level=1)
        d.add_paragraph("이상 현상을 없애는 설계 방법")
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "정규형", "조건"
        table.cell(1, 0).text, table.cell(1, 1).text = "1NF", "원자값"
        buf = io.BytesIO()
        d.save(buf)
        md = ai.extract_text(io.BytesIO(buf.getvalue()), "note.docx")
        self.assertIn("# 정규화", md)
        self.assertIn("|1NF|원자값|", md)

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "트랜잭션"
        slide.placeholders[1].text = "원자성 일관성 고립성 지속성"
        buf = io.BytesIO()
        prs.save(buf)
        md = ai.extract_text(io.BytesIO(buf.getvalue()), "slides.pptx")
        self.assertIn("트랜잭션", md)
        self.assertIn("원자성", md)

    def test_short_answer_local_match(self):
        self.assertTrue(ai.matches_short_answer("정보 처리 | information processing", "Information-Processing"))
        # ' / '는 항목 구분일 수 있어 대안으로 보지 않는다 (6개 중 하나만 써도 정답이 되면 안 된다)
        self.assertFalse(ai.matches_short_answer("Ut / Re / Mi / Fa / Sol / La", "Ut"))
        self.assertTrue(ai.matches_short_answer("Ut / Re / Mi / Fa / Sol / La", "Ut Re Mi Fa Sol La"))
        self.assertEqual(ai.answer_diff("Ut / Re / Mi / Fa / Sol / La", "Ut Le Mi Fa Sol La"), (["le"], ["re"]))
        self.assertFalse(ai.matches_short_answer("LEFT JOIN", "JOIN"))
        with patch("core.ai.complete_json") as llm:
            self.assertTrue(ai.grade_answer("q", "LEFT JOIN", "left join", question_type="short")["is_correct"])
        llm.assert_not_called()


@override_settings(AI_CREDIT_COSTS={"summary": 3, "question": 2, "grading": 1})
class TokenSavingTestCase(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="u1", password="pw")
        self.client.force_login(self.user)

    def upload(self, content=b"hello world " * 20, name="note.txt"):
        with patch("mypage.views.threading.Thread") as thread:
            res = self.client.post("/mypage/", {"title": "", "file": SimpleUploadedFile(name, content),
                                                "subject_code": "", "document_type": "lecture"})
        return res, thread

    def test_same_file_reuses_summary_without_ai(self):
        from accounts import credits
        credits.grant(self.user, 3, "test")
        res, thread = self.upload()
        doc = PersonalDocument.objects.get()
        self.assertEqual(len(doc.content_hash), 64)
        thread.assert_called_once()  # 첫 업로드는 요약 생성
        PersonalDocument.objects.filter(pk=doc.pk).update(summary="이미 만든 요약", summary_status=ProcessingStatus.DONE)
        other = User.objects.create_user(username="u2", password="pw")
        self.client.force_login(other)
        res, thread = self.upload()
        reused = PersonalDocument.objects.get(user=other)
        self.assertEqual((reused.summary, reused.summary_status), ("이미 만든 요약", ProcessingStatus.DONE))
        self.assertNotEqual(thread.call_args.kwargs.get("target").__name__, "_generate_summary_bg")

    def test_upload_blocked_when_summary_quota_used(self):
        res, thread = self.upload(b"another file content " * 10)       # 크레딧 0
        thread.assert_not_called()
        self.assertFalse(PersonalDocument.objects.exists())

    def test_ox_explanation_stored_at_generation_and_reused(self):
        doc = PersonalDocument.objects.create(user=self.user, title="d", file="d.txt", extracted_text="본문")
        q = GeneratedQuestion.objects.create(document=doc, question_type="ox", status=ProcessingStatus.PROCESSING)
        with patch("mypage.views.generate_one_question",
                   return_value={"question": "Q?", "answer": "O", "explanation": "근거 설명"}):
            _generate_question_bg(q.pk, "본문", "ox", [], self.user)
        q.refresh_from_db()
        self.assertEqual(q.feedback, "근거 설명")
        with patch("core.ai._client") as client:
            self.client.post(f"/mypage/{doc.pk}/answer/{q.pk}/", {"user_answer": "X"})
        client.assert_not_called()
        q.refresh_from_db()
        self.assertEqual((q.is_correct, q.feedback), (False, "근거 설명"))

    def test_question_generation_blocked_by_quota(self):
        doc = PersonalDocument.objects.create(user=self.user, title="d", file="d.txt", extracted_text="본문")
        with patch("mypage.views.threading.Thread") as thread:                 # 크레딧 0
            self.client.post(f"/mypage/{doc.pk}/generate/", {"question_type": "ox"})
        thread.assert_not_called()
        self.assertFalse(GeneratedQuestion.objects.exists())

    def test_question_prompt_keeps_document_in_cached_prefix(self):
        captured = {}

        def fake(task, messages, **kwargs):
            captured["messages"] = messages
            return {"question": "Q", "answer": "A"}

        with patch("core.ai.complete_json", side_effect=fake):
            ai.generate_one_question("문서 본문 XYZ", "short", ["이전 문제"], user=self.user)
        system, request = captured["messages"]
        self.assertIn("문서 본문 XYZ", system["content"])      # 매번 같은 앞부분 (캐시됨)
        self.assertNotIn("이전 문제", system["content"])        # 바뀌는 부분은 뒤로
        self.assertIn("이전 문제", request["content"])
